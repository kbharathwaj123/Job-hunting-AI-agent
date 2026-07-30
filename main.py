"""
Job Agent — main orchestrator.

Run with:  python main.py

Flow per run:
  1. Load config + init DB
  2. Open persistent browser (log in manually on first run)
  3. Search LinkedIn / Naukri / Indeed for configured roles
  4. Skip jobs already seen, jobs over the daily cap, excluded keywords
  5. Score each job against your base resume (ATS-style)
  6. Tailor resume via local LLM using the keyword gap
  7. Apply (auto-submit or stage, per config, per source)
  8. Log everything to data/applications.db
"""

import re
import yaml
import subprocess
import time
import requests as http_requests
from pathlib import Path
from docx import Document

from browser.session import get_browser_context
from data.db import init_db, already_seen, record, today_count
from ats.scorer import score_resume_against_job
from ats.company_check import verify_company
from resume.tailor import tailor_resume
from sites import linkedin
from sites.naukri_indeed import search_naukri, search_indeed, apply_naukri, apply_indeed
from sites.company_careers import apply_company_website
from resume.report import generate_pdf_report

CONFIG_PATH = Path(__file__).parent / "config" / "config.yaml"


def load_config():
    with open(CONFIG_PATH) as f:
        return yaml.safe_load(f)


def load_base_resume_text(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def ensure_ollama_running(host: str = "http://localhost:11434"):
    """Check if Ollama is running, start it if not."""
    try:
        http_requests.get(host, timeout=3)
        print("[OK] Ollama is running.")
        return True
    except Exception:
        print("[STARTING] Ollama is not running. Starting it now...")
        try:
            subprocess.Popen(
                ["ollama", "serve"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
            # Wait for it to be ready
            for i in range(15):
                time.sleep(2)
                try:
                    http_requests.get(host, timeout=2)
                    print("[OK] Ollama started successfully.")
                    return True
                except Exception:
                    pass
            print("[WARNING] Could not start Ollama. Resume tailoring will be skipped.")
            return False
        except FileNotFoundError:
            print("[WARNING] Ollama not found. Install from https://ollama.com/download")
            return False


def ensure_browser_valid(playwright, context, headless: bool):
    """Verify the browser is still open and responsive. Re-launch if not."""
    try:
        # Test if we can open and close a dummy page
        page = context.new_page()
        page.close()
        return playwright, context
    except Exception:
        print("\n[INFO] Browser was closed or crashed. Re-launching a fresh browser window...")
        try:
            context.close()
        except Exception:
            pass
        try:
            playwright.stop()
        except Exception:
            pass
        
        # Re-launch context
        from browser.session import get_browser_context
        pw, ctx = get_browser_context(headless=headless)
        ctx.set_default_timeout(30000)
        return pw, ctx


import sys

def prompt_application_mode() -> str:
    """Prompt the user in terminal or parse CLI arguments for mode selection."""
    if "--mode" in sys.argv:
        idx = sys.argv.index("--mode")
        if idx + 1 < len(sys.argv) and sys.argv[idx + 1] == "2":
            print(">> CLI Mode Selected: [2] Companies Official Website\n")
            return "company_website"
        else:
            print(">> CLI Mode Selected: [1] Online Portals\n")
            return "portals"

    print("\n===========================================================")
    print("           JOB APPLICATION MODE SELECTION                  ")
    print("===========================================================")
    print("Do you want to apply jobs through Portals or Companies Official Website?")
    print("  [1] Online Portals (LinkedIn, Naukri, Indeed)")
    print("  [2] Companies Official Website (Official Careers Pages & ATS)")
    print("===========================================================")
    try:
        choice = input("Enter choice (1 or 2, default is 1): ").strip()
        if choice == "2":
            print(">> Mode Selected: [2] Companies Official Website\n")
            return "company_website"
    except Exception:
        pass
    print(">> Mode Selected: [1] Online Portals\n")
    return "portals"


from data.history import prune_old_records, is_already_applied, add_applied_job


def main():
    cfg = load_config()
    init_db()

    # Clean history records older than 90 days (3 months)
    prune_old_records(90)

    # Ensure Ollama is running before we need it
    llm_host = cfg.get("llm", {}).get("host", "http://localhost:11434")
    ollama_ok = ensure_ollama_running(llm_host)

    # Ask user or parse CLI args for application mode
    app_mode = prompt_application_mode()

    resume_path = cfg["paths"]["base_resume"]
    if not Path(resume_path).exists():
        print(f"[WARNING] No resume found at {resume_path}. "
              f"Export your resume as .docx and place it there, then re-run.")
        return

    base_resume_text = load_base_resume_text(resume_path)
    cap = cfg["automation"]["daily_application_cap"]

    # Check CLI or config for headless background execution
    headless_mode = cfg["automation"]["headless"]
    if "--headless" in sys.argv or "--background" in sys.argv:
        headless_mode = True

    playwright, context = get_browser_context(headless=headless_mode)
    context.set_default_timeout(30000)
    reviews_page = context.new_page()

    all_jobs = []
    if cfg["sources"]["linkedin"]["enabled"]:
        try:
            print("[SEARCHING] LinkedIn...")
            playwright, context = ensure_browser_valid(playwright, context, cfg["automation"]["headless"])
            jobs = linkedin.search_jobs(
                context, cfg["job_criteria"]["roles"],
                cfg["job_criteria"]["locations"], cfg["job_criteria"]["wfh_preference"],
            )
            all_jobs += jobs
            print(f"  -> Found {len(jobs)} jobs on LinkedIn")
        except Exception as e:
            print(f"  [ERROR] LinkedIn search failed: {e}")

    if cfg["sources"]["naukri"]["enabled"]:
        try:
            print("[SEARCHING] Naukri...")
            playwright, context = ensure_browser_valid(playwright, context, cfg["automation"]["headless"])
            jobs = search_naukri(context, cfg["job_criteria"]["roles"], cfg["job_criteria"]["locations"])
            all_jobs += jobs
            print(f"  -> Found {len(jobs)} jobs on Naukri")
        except Exception as e:
            print(f"  [ERROR] Naukri search failed: {e}")

    if cfg["sources"]["indeed"]["enabled"]:
        try:
            print("[SEARCHING] Indeed...")
            playwright, context = ensure_browser_valid(playwright, context, cfg["automation"]["headless"])
            jobs = search_indeed(context, cfg["job_criteria"]["roles"], cfg["job_criteria"]["locations"])
            all_jobs += jobs
            print(f"  -> Found {len(jobs)} jobs on Indeed")
        except Exception as e:
            print(f"  [ERROR] Indeed search failed: {e}")

    print(f"\nFound {len(all_jobs)} candidate postings total.")

    processed_jobs = []

    for i, job in enumerate(all_jobs, 1):
        if today_count() >= cap:
            print("Daily application cap reached. Stopping.")
            break
        if already_seen(job["url"]):
            continue

        # 1ST CHANGE CONDITION: Check 3-Month Application History Deduplication
        if is_already_applied(job["company"], job["title"]):
            print(f"  [SKIP - ALREADY APPLIED 🔁] Already applied to '{job['title']}' @ '{job['company']}' in the last 3 months.")
            record(job["source"], job["title"], job["company"], job["url"], 0, "already_applied")
            processed_jobs.append({
                "company": job["company"], "title": job["title"], "source": job["source"],
                "ats_score": 0, "status": "already_applied", "status_reason": "Already applied in last 3 months",
                "location": job.get("location", "Not Specified"), "wfh": "Any", "salary": "Not Specified",
                "company_email": "Not Listed"
            })
            continue

        # Check Excluded Companies (Current / Previous Employer Protection)
        excluded_companies = [c.strip().lower() for c in cfg["job_criteria"].get("exclude_companies", []) if c.strip()]
        curr_comp = cfg.get("profile_answers", {}).get("current_company", "").strip().lower()
        if curr_comp and curr_comp not in excluded_companies:
            excluded_companies.append(curr_comp)

        job_company_clean = job["company"].strip().lower()
        # Filter out generic template placeholders
        active_exclusions = [exc for exc in excluded_companies if exc and exc not in ("your current company name", "previous employer name", "example company")]

        if any(exc in job_company_clean or job_company_clean in exc for exc in active_exclusions):
            print(f"  [SKIP - CURRENT EMPLOYER 🛑] Skipping '{job['company']}' as it matches your current employer.")
            record(job["source"], job["title"], job["company"], job["url"], 0, "skipped_current_employer")
            processed_jobs.append({
                "company": job["company"], "title": job["title"], "source": job["source"],
                "ats_score": 0, "status": "skipped", "status_reason": f"Current employer ({job['company']}) - excluded",
                "location": job.get("location", "Not Specified"), "wfh": "Any", "salary": "Not Specified",
                "company_email": "Not Listed"
            })
            continue

        if any(bad.lower() in job["title"].lower() for bad in cfg["job_criteria"]["exclude_keywords"]):
            record(job["source"], job["title"], job["company"], job["url"], 0, "skipped")
            processed_jobs.append({
                "company": job["company"], "title": job["title"], "source": job["source"],
                "ats_score": 0, "status": "skipped", "status_reason": "Title contains excluded keyword",
                "location": job.get("location", "Not Specified"),
                "wfh": cfg["job_criteria"]["wfh_preference"], "salary": "Not Specified",
                "company_email": "Not Listed"
            })
            continue

        try:
            print(f"\n[{i}/{len(all_jobs)}] Processing: {job['title']} @ {job['company']} ({job['source']})")

            playwright, context = ensure_browser_valid(playwright, context, cfg["automation"]["headless"])

            try:
                reviews_page.goto("about:blank")
            except Exception:
                reviews_page = context.new_page()

            # Company verification — threshold rating >= 3.5 & reviews >= 80
            if cfg["job_criteria"].get("verify_company", False):
                llm_cfg = cfg.get("llm", {})
                min_r = cfg["job_criteria"].get("min_company_rating", 3.5)
                min_rev = cfg["job_criteria"].get("min_company_reviews", 80)
                check = verify_company(
                    job["company"],
                    page=reviews_page,
                    job_description=job.get("description", ""),
                    min_rating=min_r,
                    min_reviews=min_rev,
                    ollama_host=llm_cfg.get("host", "http://localhost:11434"),
                    model=llm_cfg.get("model", "llama3.1:8b"),
                )
                if check["verdict"] == "suspicious":
                    print(f"  [SKIPPED] Company looks suspicious: {check['reason']}")
                    record(job["source"], job["title"], job["company"], job["url"], 0, "suspicious")
                    processed_jobs.append({
                        "company": job["company"], "title": job["title"], "source": job["source"],
                        "ats_score": 0, "status": "suspicious", "status_reason": check['reason'],
                        "location": job.get("location", "Not Specified"),
                        "wfh": cfg["job_criteria"]["wfh_preference"], "salary": "Not Specified",
                        "company_email": "Not Listed"
                    })
                    continue
                elif check["verdict"] == "skip_low_rating":
                    print(f"  [SKIPPED] Low company rating/reviews: {check['reason']}")
                    record(job["source"], job["title"], job["company"], job["url"], 0, "skipped_low_rating")
                    processed_jobs.append({
                        "company": job["company"], "title": job["title"], "source": job["source"],
                        "ats_score": 0, "status": "skipped_low_rating", "status_reason": check['reason'],
                        "location": job.get("location", "Not Specified"),
                        "wfh": cfg["job_criteria"]["wfh_preference"], "salary": "Not Specified",
                        "company_email": "Not Listed"
                    })
                    continue
                elif check["verdict"] == "legit":
                    print(f"  [VERIFIED] {check['reason']}")

            job_description = job.get("description", job["title"])
            
            email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', job_description)
            company_email = email_match.group(0) if email_match else "Not Listed"

            score, matched, missing = score_resume_against_job(base_resume_text, job_description)

            print(f"  ATS Score: {score}% | Matched: {len(matched)} | Missing: {len(missing)}")

            # 3RD CHANGE CONDITION: ATS Score Threshold (>= 55%)
            min_ats = cfg["job_criteria"].get("min_ats_score", 55)
            if score < min_ats:
                reason_ats = f"ATS match {score}% < {min_ats}% min threshold"
                print(f"  [SKIPPED] {reason_ats}")
                record(job["source"], job["title"], job["company"], job["url"], score, "low_ats_score")
                processed_jobs.append({
                    "company": job["company"], "title": job["title"], "source": job["source"],
                    "ats_score": score, "status": "low_ats_score", "status_reason": reason_ats,
                    "location": job.get("location", "Not Specified"), "wfh": "Any", "salary": "Not Specified",
                    "company_email": company_email
                })
                continue

            tailored = tailor_resume(base_resume_text, job_description, missing)

            source_cfg = cfg["sources"].get(job["source"], {})
            auto_submit = source_cfg.get("auto_submit", False)

            profile_answers = cfg.get("profile_answers", {})

            if app_mode == "company_website":
                print(f"  [MODE 2: COMPANY WEBSITE] Navigating to official careers portal for '{job['company']}'...")
                result = apply_company_website(
                    context, job["company"], job["title"], job["url"],
                    base_resume_text, profile_answers, resume_path, auto_submit
                )
                source_label = "Company Portal"
            else:
                source_label = job["source"].capitalize()
                if job["source"] == "linkedin":
                    result = linkedin.easy_apply(context, job["url"], base_resume_text, profile_answers, tailored.get("summary", ""), auto_submit, company_name=job["company"])
                elif job["source"] == "naukri":
                    result = apply_naukri(context, job["url"], base_resume_text, profile_answers,
                                           resume_path, auto_submit, company_name=job["company"])
                elif job["source"] == "indeed":
                    pdf_path = cfg["paths"].get("base_resume_pdf", "")
                    result = apply_indeed(context, job["url"], base_resume_text, profile_answers,
                                           pdf_path, auto_submit, company_name=job["company"])
                else:
                    result = {"status": "staged", "reason": f"{job['source']} apply flow not yet built"}

            shot_path = result.get("screenshot", "")
            record(job["source"], job["title"], job["company"], job["url"], score, result["status"], screenshot=shot_path)
            print(f"  [{result['status'].upper()}] {job['title']} @ {job['company']} (score={score}%)")
            
            # If successfully applied or submitted, add to 3-month history file
            if result["status"] in ("applied", "submitted"):
                add_applied_job(job)

            processed_jobs.append({
                "company": job["company"], "title": job["title"], "source": source_label,
                "ats_score": score, "status": result["status"], "status_reason": result.get("reason", "Successfully Processed"),
                "location": job.get("location", "Hyderabad / Pune / Global"),
                "wfh": "On-site / Hybrid / Remote", "salary": "Not Specified",
                "company_email": company_email, "screenshot": shot_path
            })
        except Exception as e:
            print(f"  [ERROR] Failed to process {job['title']}: {e}")
            record(job["source"], job["title"], job["company"], job["url"], 0, "error")
            processed_jobs.append({
                "company": job["company"], "title": job["title"], "source": job["source"],
                "ats_score": 0, "status": "error", "location": "Not Specified",
                "wfh": cfg["job_criteria"]["wfh_preference"], "salary": "Not Specified",
                "company_email": "Not Listed"
            })

    try:
        reviews_page.close()
    except Exception:
        pass
    try:
        context.close()
    except Exception:
        pass
    try:
        playwright.stop()
    except Exception:
        pass

    if processed_jobs:
        generate_pdf_report(processed_jobs, "C:/Users/HP/Downloads/JobsApplied.pdf")


if __name__ == "__main__":
    main()
